import io
from openpyxl import Workbook
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors

from services.estado_resultados_service import generar_estado_resultados
from services.balance_service import calcular_balance_general
from services.consolidacion_service import consolidar_informacion_contable


# ---------------------------------------------------------------
# Helpers de formato
# ---------------------------------------------------------------
def _monto(valor):
    try:
        return f"${float(valor):,.2f}"
    except (TypeError, ValueError):
        return "$0.00"


def _encabezado_reporte(titulo, subtitulo=None):
    styles = getSampleStyleSheet()
    elementos = []
    elementos.append(Paragraph(f"<b>{titulo}</b>", styles['Title']))
    if subtitulo:
        elementos.append(Paragraph(subtitulo, styles['Normal']))
        elementos.append(Spacer(1, 6))
    elementos.append(Spacer(1, 12))
    return styles, elementos


def _construir_tabla(datos_tabla, ancho_col=None, resaltar_ultima_fila=True):
    """Crea una tabla con estilo corporativo."""
    t = Table(datos_tabla, colWidths=ancho_col)
    estilo = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2C3E50')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
    ]
    if resaltar_ultima_fila:
        estilo.append(('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#ECF0F1')))
        estilo.append(('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'))
    t.setStyle(TableStyle(estilo))
    return t


# ---------------------------------------------------------------
# Estado de Resultados
# ---------------------------------------------------------------
def exportar_estado_resultados_pdf(datos_reporte, periodo_txt=""):
    """Genera un archivo PDF en memoria del Estado de Resultados."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    elementos = []

    elementos.append(Paragraph("<b>ESTADO DE RESULTADOS CONSOLIDADO</b>", styles['Title']))
    if periodo_txt:
        elementos.append(Paragraph(f"Periodo: {periodo_txt}", styles['Normal']))
    elementos.append(Spacer(1, 12))

    tabla_datos = [
        ['Rubro Contable', 'Monto ($)'],
        ['Ingresos Totals', _monto(datos_reporte['ingresos_totales'])],
        ['Costo de Ventas', _monto(datos_reporte['costo_ventas'])],
        ['Utilidad Bruta', _monto(datos_reporte['utilidad_bruta'])],
        ['Gastos Operativos (Nómina + Deducibles + Pérdidas)', _monto(datos_reporte['gastos_operativos'])],
        ['UTILIDAD NETA', _monto(datos_reporte['utilidad_neta'])]
    ]
    elementos.append(_construir_tabla(tabla_datos, ancho_col=[300, 150]))
    doc.build(elementos)
    buffer.seek(0)
    return buffer


def exportar_estado_resultados_excel(datos_reporte, periodo_txt=""):
    """Genera una hoja de cálculo Excel en memoria del Estado de Resultados."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Estado de Resultados"

    ws.append(["Estado de Resultados Consolidado"])
    if periodo_txt:
        ws.append(["Periodo", periodo_txt])
    ws.append([])
    ws.append(["Rubro Contable", "Monto ($)"])
    ws.append(["Ingresos Totals", datos_reporte['ingresos_totales']])
    ws.append(["Costo de Ventas", datos_reporte['costo_ventas']])
    ws.append(["Utilidad Bruta", datos_reporte['utilidad_bruta']])
    ws.append(["Gastos Operativos", datos_reporte['gastos_operativos']])
    ws.append(["UTILIDAD NETA", datos_reporte['utilidad_neta']])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ---------------------------------------------------------------
# Balance General
# ---------------------------------------------------------------
def exportar_balance_general_pdf(datos_balance, periodo_txt=""):
    """Genera un archivo PDF en memoria del Balance General."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    elementos = []

    elementos.append(Paragraph("<b>BALANCE GENERAL</b>", styles['Title']))
    if periodo_txt:
        elementos.append(Paragraph(f"Periodo: {periodo_txt}", styles['Normal']))
    elementos.append(Spacer(1, 12))

    tabla_datos = [
        ['Estructura Financiera', 'Valor ($)'],
        ['ACTIVOS', ''],
        ['Efectivo / Flujo de Caja Disponible', _monto(datos_balance['activos']['efectivo_caja'])],
        ['Inventario de Mercancía (a costo)', _monto(datos_balance['activos']['inventario'])],
        ['Total Activos', _monto(datos_balance['activos']['total_activos'])],
        ['PASIVOS', ''],
        ['Pasivos Laborales (Retenciones)', _monto(datos_balance['pasivos']['pasivos_laborales'])],
        ['Total Pasivos', _monto(datos_balance['pasivos']['total_pasivos'])],
        ['PATRIMONIO', ''],
        ['Utilidad del Ejercicio', _monto(datos_balance['patrimonio']['utilidad_ejercicio'])],
        ['Total Patrimonio', _monto(datos_balance['patrimonio']['total_patrimonio'])]
    ]
    elementos.append(_construir_tabla(tabla_datos, ancho_col=[340, 140]))
    doc.build(elementos)
    buffer.seek(0)
    return buffer


def exportar_balance_general_excel(datos_balance, periodo_txt=""):
    """Genera una hoja de cálculo Excel en memoria del Balance General."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Balance General"

    ws.append(["Balance General"])
    if periodo_txt:
        ws.append(["Periodo", periodo_txt])
    ws.append([])
    ws.append(["Estructura Financiera", "Valor ($)"])
    ws.append(["ACTIVOS", ""])
    ws.append(["Efectivo / Flujo de Caja Disponible", datos_balance['activos']['efectivo_caja']])
    ws.append(["Inventario de Mercancía (a costo)", datos_balance['activos']['inventario']])
    ws.append(["Total Activos", datos_balance['activos']['total_activos']])
    ws.append(["PASIVOS", ""])
    ws.append(["Pasivos Laborales (Retenciones)", datos_balance['pasivos']['pasivos_laborales']])
    ws.append(["Total Pasivos", datos_balance['pasivos']['total_pasivos']])
    ws.append(["PATRIMONIO", ""])
    ws.append(["Utilidad del Ejercicio", datos_balance['patrimonio']['utilidad_ejercicio']])
    ws.append(["Total Patrimonio", datos_balance['patrimonio']['total_patrimonio']])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ---------------------------------------------------------------
# Consolidado
# ---------------------------------------------------------------
def exportar_consolidado_pdf(datos_consol, periodo_txt=""):
    """Genera un PDF con la información consolidada de inventario, nómina y gastos."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    elementos = []

    elementos.append(Paragraph("<b>CONSOLIDACIÓN DE INFORMACIÓN CONTABLE</b>", styles['Title']))
    if periodo_txt:
        elementos.append(Paragraph(f"Periodo: {periodo_txt}", styles['Normal']))
    elementos.append(Spacer(1, 12))

    tabla_datos = [
        ['Concepto', 'Valor ($)'],
        ['Ingresos Totals', _monto(datos_consol['ingresos_totals'])],
        ['Costo de Ventas', _monto(datos_consol['costo_ventas'])],
        ['Compras de Inventario', _monto(datos_consol['compras_inventario'])],
        ['Pérdidas de Inventario', _monto(datos_consol['perdidas_inventario'])],
        ['Costo de Nómina', _monto(datos_consol['costo_nomina'])],
        ['Deducciones Nómina', _monto(datos_consol['deducciones_nomina'])],
        ['Gastos Deducibles', _monto(datos_consol['gastos_deducibles'])],
        ['Valor Inventario (final)', _monto(datos_consol['valor_inventario'])],
        ['Valor Inventario (inicial)', _monto(datos_consol['valor_inventario_inicial'])]
    ]
    elementos.append(_construir_tabla(tabla_datos, ancho_col=[300, 150]))
    doc.build(elementos)
    buffer.seek(0)
    return buffer


def exportar_consolidado_excel(datos_consol, periodo_txt=""):
    """Genera un Excel con la información consolidada."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado Contable"

    ws.append(["Consolidación de Información Contable"])
    if periodo_txt:
        ws.append(["Periodo", periodo_txt])
    ws.append([])
    ws.append(["Concepto", "Valor ($)"])
    filas = [
        ("Ingresos Totals", datos_consol['ingresos_totals']),
        ("Costo de Ventas", datos_consol['costo_ventas']),
        ("Compras de Inventario", datos_consol['compras_inventario']),
        ("Pérdidas de Inventario", datos_consol['perdidas_inventario']),
        ("Costo de Nómina", datos_consol['costo_nomina']),
        ("Deducciones Nómina", datos_consol['deducciones_nomina']),
        ("Gastos Deducibles", datos_consol['gastos_deducibles']),
        ("Valor Inventario (final)", datos_consol['valor_inventario']),
        ("Valor Inventario (inicial)", datos_consol['valor_inventario_inicial']),
    ]
    for concepto, valor in filas:
        ws.append([concepto, valor])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ---------------------------------------------------------------
# Funciones de conveniencia para rutas (obtener datos + periodo)
# ---------------------------------------------------------------
def obtener_estado_resultados_con_periodo(fecha_inicio=None, fecha_fin=None):
    periodo_txt = ""
    if fecha_inicio and fecha_fin:
        periodo_txt = f"{fecha_inicio} al {fecha_fin}"
    datos = generar_estado_resultados(fecha_inicio, fecha_fin)
    return datos, periodo_txt


def obtener_balance_con_periodo(fecha_inicio=None, fecha_fin=None):
    periodo_txt = ""
    if fecha_inicio and fecha_fin:
        periodo_txt = f"{fecha_inicio} al {fecha_fin}"
    datos = calcular_balance_general(fecha_inicio, fecha_fin)
    return datos, periodo_txt


def obtener_consolidado_con_periodo(fecha_inicio=None, fecha_fin=None):
    periodo_txt = ""
    if fecha_inicio and fecha_fin:
        periodo_txt = f"{fecha_inicio} al {fecha_fin}"
    datos = consolidar_informacion_contable(fecha_inicio, fecha_fin)
    return datos, periodo_txt
